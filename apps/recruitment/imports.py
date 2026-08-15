"""استيراد جماعي لبيانات التوظيف من CSV / Excel / JSON / SQL (v4).

الصيغ المدعومة:
- CSV / XLSX: أول ورقة، الترويسة = أسماء الأعمدة (entity من الصفحة).
- JSON: قائمة كائنات أو كائن واحد (أو مفتاح records/postings/candidates).
- SQL: جمل INSERT INTO (قد تتضمن جداول إعلانات ومرشحين معًا) تُكتشف تلقائيًا.

يعالج الصفوف واحدًا واحدًا:
- كل صف يُتحقق منه مستقلاً؛ تُسجَّل الأخطاء ولا يُتوقف الملف عند أولها.
- قواعد العمل BR-REC-001..005 تبقى سارية (لا يُضاف مرشح لإعلان مغلق...).
- الأعمدة الاختيارية فارغة أو مفقودة تُهمل.
"""

import csv
import io
import json
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from django.core.exceptions import ValidationError as DjangoValidationError
from django.core.validators import validate_email
from django.utils.translation import gettext as _
from openpyxl import load_workbook

from apps.org.models import Branch, Department

from .models import Candidate, JobPosting
from .services import RecruitmentError, add_candidate


class ImportFileError(Exception):
    """خطأ على مستوى الملف (امتداد أو بنية غير صالحة)."""


def _clean(value):
    """يحوّل خلية Excel/CSV إلى نص نظيف مع الحفاظ على التواريخ والأرقام."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_table(file):
    """يقرأ CSV/XLSX/JSON ويعيد قائمة قواميس (مفاتيح = أعمدة الترويسة)."""
    name = (getattr(file, "name", "") or "").lower()
    if not name.endswith((".csv", ".xlsx", ".json")):
        raise ImportFileError(_("امتداد غير مدعوم — استخدم CSV أو XLSX أو JSON"))
    if name.endswith(".json"):
        return _read_json_records(file)
    raw = file.read()
    if name.endswith(".csv"):
        reader = csv.DictReader(io.StringIO(raw.decode("utf-8-sig")))
        return [dict(row) for row in reader]
    workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    header = [str(head).strip() for head in (next(rows, None) or [])]
    records = []
    for values in rows:
        if values is None or all(value is None for value in values):
            continue
        records.append({head: _clean(value) for head, value in zip(header, values) if head})
    return records


def _read_json_records(file):
    """يقرأ ملف JSON: قائمة كائنات، أو كائن واحد، أو مفتاح records/postings/candidates."""
    raw = file.read().decode("utf-8-sig", errors="replace")
    try:
        data = json.loads(raw)
    except ValueError as exc:
        raise ImportFileError(_("ملف JSON غير صالح: %(s)s") % {"s": exc})
    if isinstance(data, dict):
        for key in ("records", "postings", "candidates"):
            if isinstance(data.get(key), list):
                data = data[key]
                break
        else:
            data = [data]
    if not isinstance(data, list):
        raise ImportFileError(_("JSON يجب أن يكون قائمة كائنات"))
    if not all(isinstance(item, dict) for item in data):
        raise ImportFileError(_("جميع عناصر JSON يجب أن تكون كائنات"))
    return data


# ---- دعم ملفات SQL (INSERT INTO ... VALUES) --------------------------------

_SQL_TABLE_ALIASES = {
    "postings": {
        "posting", "postings",
        "jobposting", "jobpostings",
        "recruitment_jobposting", "recruitment_jobpostings",
    },
    "candidates": {
        "candidate", "candidates",
        "recruitment_candidate", "recruitment_candidates",
    },
}

_INSERT_RE = re.compile(
    r"INSERT\s+INTO\s+[`\"\[]?([\w.]+)[`\"\]]?\s*(?:\(([^)]*)\))?\s*VALUES\s*",
    re.IGNORECASE,
)


def _sql_entity(table_name):
    base = table_name.strip().split(".")[-1].strip().lower()
    for entity, aliases in _SQL_TABLE_ALIASES.items():
        if base in aliases:
            return entity
    return None


def _split_top_level(text, sep=","):
    """يقسّم نصًا على فاصل مع احترام الأقواس والاقتباسات."""
    parts, buf, depth, quote = [], [], 0, None
    for ch in text:
        if quote:
            buf.append(ch)
            if ch == quote:
                quote = None
        elif ch in ("'", '"'):
            quote = ch
            buf.append(ch)
        elif ch in "([":
            depth += 1
            buf.append(ch)
        elif ch in ")]":
            depth -= 1
            buf.append(ch)
        elif ch == sep and depth == 0:
            parts.append("".join(buf).strip())
            buf = []
        else:
            buf.append(ch)
    if buf:
        parts.append("".join(buf).strip())
    return parts


def _sql_row_groups(text):
    """يستخرج مجموعات الصفوف (...) من نص بعد VALUES (يوقف عند ; أو نهاية النص)."""
    groups = []
    depth, quote = 0, None
    cur = []
    for ch in text:
        if quote:
            cur.append(ch)
            if ch == quote:
                quote = None
        elif ch in ("'", '"'):
            quote = ch
            cur.append(ch)
        elif ch == "(":
            depth += 1
            if depth == 1:
                cur = []
            cur.append(ch)
        elif ch == ")":
            depth -= 1
            cur.append(ch)
            if depth == 0:
                groups.append(_split_top_level("".join(cur)[1:-1]))
        elif ch == ";":
            break
        else:
            if depth > 0:
                cur.append(ch)
    return groups


def _sql_value(value):
    value = value.strip()
    if not value or value.upper() == "NULL":
        return ""
    if value.startswith("'") and value.endswith("'"):
        return value[1:-1].replace("''", "'")
    if value.startswith('"') and value.endswith('"'):
        return value[1:-1].replace('""', '"')
    return value


def read_sql_tables(file):
    """يحلل ملف SQL ويعيد {entity: [records]} (جداول إعلانات ومرشحين معًا).

    يتطلب أعمدة صريحة بعد INSERT INTO (لا يدعم VALUES بدون أعمدة).
    """
    text = file.read().decode("utf-8-sig", errors="replace")
    result = {"postings": [], "candidates": []}
    found = False
    matches = list(_INSERT_RE.finditer(text))
    for index, match in enumerate(matches):
        entity = _sql_entity(match.group(1))
        if not entity:
            continue
        columns = [col.strip().strip("`\"'[]") for col in (match.group(2) or "").split(",") if col.strip()]
        if not columns:
            raise ImportFileError(_("جمل INSERT تتطلب أعمدة صريحة بعد اسم الجدول"))
        start = match.end()
        end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        for group in _sql_row_groups(text[start:end]):
            values = [_sql_value(value) for value in group]
            if len(values) != len(columns):
                raise ImportFileError(
                    _("عدد القيم (%(n)d) لا يطابق عدد الأعمدة (%(c)d)")
                    % {"n": len(values), "c": len(columns)}
                )
            result[entity].append(dict(zip(columns, values)))
            found = True
    if not found:
        raise ImportFileError(_("لم يُعثر على جداول مدعومة (postings/candidates) في ملف SQL"))
    return result


def _as_date(value):
    value = (value or "").strip()
    if not value:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(value, fmt).date()
        except ValueError:
            continue
    raise ValueError(_("تاريخ غير صالح: %(v)s") % {"v": value})


def _as_int(value, default):
    value = (value or "").strip()
    if not value:
        return default
    try:
        return int(Decimal(value))
    except (InvalidOperation, ValueError):
        raise ValueError(_("رقم غير صالح: %(v)s") % {"v": value})


def _lookup_department(code):
    code = (code or "").strip()
    if not code:
        return None
    return Department.objects.filter(code=code).first()


def _lookup_branch(code):
    code = (code or "").strip()
    if not code:
        return None
    return Branch.objects.filter(code=code).first()


def _normalize_records(records):
    """يحوّل قيم JSON/XLSX الأصلية (int/float/date/None/...) إلى نصوص نظيفة."""
    return [{key: _clean(value) for key, value in record.items()} for record in records]


def import_postings(records, user=None):
    """يستورد إعلانات وظائف جديدة. يعيد (عدد المُنشأ، قائمة أخطاء صفوف).

    الأعمدة: code, title_ar, title_fr, title_en, department_code, branch_code,
    employment_type, openings_count, status, publish_date, close_date,
    requirements, description.
    """
    records = _normalize_records(records)
    created = 0
    errors = []
    for index, record in enumerate(records, start=2):
        code = (record.get("code") or "").strip()
        title_ar = (record.get("title_ar") or "").strip()
        row_errors = []
        if not code:
            row_errors.append(_("الرمز (code) مطلوب"))
        if not title_ar:
            row_errors.append(_("العنوان (title_ar) مطلوب"))
        if code and JobPosting.objects.filter(code=code).exists():
            row_errors.append(_("رمز مكرر: %(s)s") % {"s": code})

        status = (record.get("status") or JobPosting.Status.DRAFT).strip() or JobPosting.Status.DRAFT
        if status not in JobPosting.Status.values:
            row_errors.append(_("حالة غير معروفة: %(s)s") % {"s": status})
        employment = (record.get("employment_type") or JobPosting.EmploymentType.FULL_TIME).strip()
        if employment not in JobPosting.EmploymentType.values:
            row_errors.append(_("نوع توظيف غير معروف: %(s)s") % {"s": employment})

        department = _lookup_department(record.get("department_code"))
        if record.get("department_code") and department is None:
            row_errors.append(_("قسم غير موجود: %(s)s") % {"s": record.get("department_code")})
        branch = _lookup_branch(record.get("branch_code"))
        if record.get("branch_code") and branch is None:
            row_errors.append(_("فرع غير موجود: %(s)s") % {"s": record.get("branch_code")})

        try:
            publish_date = _as_date(record.get("publish_date"))
            close_date = _as_date(record.get("close_date"))
            openings_count = _as_int(record.get("openings_count"), 1)
        except ValueError as exc:
            row_errors.append(str(exc))

        if row_errors:
            errors.append({"row": index, "errors": row_errors})
            continue

        JobPosting.objects.create(
            code=code,
            title_ar=title_ar,
            title_fr=record.get("title_fr", "") or "",
            title_en=record.get("title_en", "") or "",
            department=department,
            branch=branch,
            employment_type=employment,
            openings_count=openings_count,
            requirements=record.get("requirements", "") or "",
            description=record.get("description", "") or "",
            status=status,
            publish_date=publish_date,
            close_date=close_date,
            created_by=user,
            updated_by=user,
        )
        created += 1
    return created, errors


def import_candidates(records, user=None):
    """يستورد مرشحين جدد (ربط اختياري بإعلان عبر posting_code).

    الأعمدة: first_name_ar, last_name_ar, email, phone, source, posting_code,
    status, first_name_fr, last_name_fr, first_name_en, last_name_en,
    applied_date, notes.
    """
    records = _normalize_records(records)
    created = 0
    errors = []
    for index, record in enumerate(records, start=2):
        first_name_ar = (record.get("first_name_ar") or "").strip()
        last_name_ar = (record.get("last_name_ar") or "").strip()
        email = (record.get("email") or "").strip()
        row_errors = []
        if not first_name_ar:
            row_errors.append(_("الاسم (first_name_ar) مطلوب"))
        if not last_name_ar:
            row_errors.append(_("اللقب (last_name_ar) مطلوب"))
        if not email:
            row_errors.append(_("البريد (email) مطلوب"))
        else:
            try:
                validate_email(email)
            except DjangoValidationError:
                row_errors.append(_("بريد غير صالح: %(v)s") % {"v": email})

        posting = None
        posting_code = (record.get("posting_code") or "").strip()
        if posting_code:
            posting = JobPosting.objects.filter(code=posting_code).first()
            if posting is None:
                row_errors.append(_("إعلان غير موجود: %(s)s") % {"s": posting_code})

        status = (record.get("status") or "").strip()
        if status and status not in Candidate.Status.values:
            row_errors.append(_("حالة غير معروفة: %(s)s") % {"s": status})

        try:
            applied_date = _as_date(record.get("applied_date"))
        except ValueError as exc:
            row_errors.append(str(exc))

        if row_errors:
            errors.append({"row": index, "errors": row_errors})
            continue

        extra = {}
        for field in (
            "first_name_fr",
            "last_name_fr",
            "first_name_en",
            "last_name_en",
            "phone",
            "source",
            "notes",
        ):
            if record.get(field):
                extra[field] = record.get(field)
        if applied_date:
            extra["applied_date"] = applied_date
        if status:
            extra["status"] = status
        try:
            add_candidate(posting, email=email, first_name_ar=first_name_ar, last_name_ar=last_name_ar, user=user, **extra)
        except RecruitmentError as exc:
            errors.append({"row": index, "errors": [str(exc)]})
            continue
        created += 1
    return created, errors
