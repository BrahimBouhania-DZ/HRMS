"""اختبارات التقارير الأساسية + التصدير (S5)."""

import csv
import datetime
import io
from decimal import Decimal
from pathlib import Path

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse

from apps.auth_app.models import Permission
from apps.attendance.models import AttendanceDay
from apps.employees.models import Employee
from apps.leave.models import LeaveBalance, LeaveRequest, LeaveType
from apps.org.models import Branch, Department
from apps.reports.scheduling import run_scheduled_reports

User = get_user_model()


class ReportTests(TestCase):
    """تغطية REP-01,10,12,13,20,21 + CSV/XLSX + النطاقات."""

    @classmethod
    def setUpTestData(cls):
        Permission.objects.create(code="reports.view", module="core", name_ar="التقارير")
        Permission.objects.create(code="reports.export", module="core", name_ar="تصدير")
        cls.admin = User.objects.create_superuser(username="rep_admin", password="pass")
        cls.emp_user = User.objects.create_user(username="rep_emp", password="pass")
        cls.branch = Branch.objects.create(code="BR-R", name_ar="فرع R")
        cls.dept = Department.objects.create(code="DEP-R", name_ar="قسم R", branch=cls.branch)
        cls.emp = Employee.objects.create(
            employee_code="R-001", first_name_ar="رأفت", last_name_ar="ر",
            branch=cls.branch, department=cls.dept, user=cls.emp_user,
            employment_status=Employee.EmploymentStatus.ACTIVE, phone="0550",
        )
        cls.emp2 = Employee.objects.create(
            employee_code="R-002", first_name_ar="رانية", last_name_ar="ر",
            branch=cls.branch, department=cls.dept,
            employment_status=Employee.EmploymentStatus.ACTIVE,
        )
        cls.today = datetime.date(2026, 8, 9)
        AttendanceDay.objects.create(employee=cls.emp, work_date=cls.today,
                                     state=AttendanceDay.State.PRESENT, late_minutes=15,
                                     worked_minutes=480)
        AttendanceDay.objects.create(employee=cls.emp2, work_date=cls.today,
                                     state=AttendanceDay.State.ABSENT)
        cls.lt = LeaveType.objects.create(code="rep-lt", name_ar="سنوية", days_per_year=30)
        LeaveBalance.objects.create(employee=cls.emp, leave_type=cls.lt, year=2026,
                                    granted=30, used=4)
        LeaveRequest.objects.create(employee=cls.emp, leave_type=cls.lt,
                                    from_date=datetime.date(2026, 9, 1),
                                    to_date=datetime.date(2026, 9, 3), days=3,
                                    status=LeaveRequest.Status.PENDING)

    def setUp(self):
        self.client.force_login(self.admin)

    def _get(self, name, params=None):
        return self.client.get(reverse(name), params or {})

    def test_index_and_all_reports_200(self):
        for name in ["reports:index", "reports:rep01", "reports:rep10",
                     "reports:rep12", "reports:rep13", "reports:rep20", "reports:rep21"]:
            self.assertEqual(self._get(name).status_code, 200, name)

    def test_employee_without_perm_forbidden(self):
        self.client.force_login(self.emp_user)
        self.assertEqual(self.client.get(reverse("reports:index")).status_code, 403)

    def test_rep01_lists_employees(self):
        resp = self._get("reports:rep01")
        self.assertContains(resp, "R-001")
        self.assertContains(resp, "رأفت")

    def test_rep10_filters_today(self):
        resp = self._get("reports:rep10", {"work_date": "2026-08-09"})
        self.assertContains(resp, "15")
        self.assertContains(resp, "480")

    def test_rep12_absence_counts(self):
        resp = self._get("reports:rep12", {"from_date": "2026-08-01", "to_date": "2026-08-31"})
        self.assertContains(resp, "R-002")

    def test_rep13_lateness_sums(self):
        resp = self._get("reports:rep13", {"from_date": "2026-08-01", "to_date": "2026-08-31"})
        self.assertContains(resp, "15")

    def test_rep20_balances(self):
        resp = self._get("reports:rep20", {"year": "2026"})
        self.assertContains(resp, "30")

    def test_rep21_requests(self):
        resp = self._get("reports:rep21")
        self.assertContains(resp, "قيد الانتظار")

    def test_export_csv(self):
        resp = self.client.get(reverse("reports:export", args=["rep01", "csv"]))
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp["Content-Type"], "text/csv; charset=utf-8")
        rows = list(csv.reader(io.StringIO(resp.content.decode("utf-8"))))
        self.assertEqual(rows[0], ["الرمز", "الاسم", "القسم", "المنصب", "الفرع", "الحالة", "الهاتف"])
        self.assertEqual(rows[1][0], "R-001")

    def test_export_xlsx(self):
        resp = self.client.get(reverse("reports:export", args=["rep01", "xlsx"]))
        self.assertEqual(resp.status_code, 200)
        self.assertIn("spreadsheetml", resp["Content-Type"])
        self.assertGreater(len(resp.content), 100)

    def test_export_pdf(self):
        resp = self.client.get(reverse("reports:export", args=["rep01", "pdf"]))
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp["Content-Type"], "application/pdf")
        self.assertTrue(resp.content.startswith(b"%PDF"))

    def test_export_requires_perm(self):
        self.client.force_login(self.emp_user)
        resp = self.client.get(reverse("reports:export", args=["rep01", "csv"]))
        self.assertEqual(resp.status_code, 403)


class ScanReportTests(TestCase):
    """REP-16/17 — سجل المسحات والمرفوضات."""

    @classmethod
    def setUpTestData(cls):
        from apps.attendance.models import AttendanceDay, AttendanceScan
        from apps.devices.models import QrDevice

        Permission.objects.create(code="reports.view", module="core", name_ar="التقارير")
        Permission.objects.create(code="reports.export", module="core", name_ar="تصدير")
        cls.admin = User.objects.create_superuser(username="scan_admin", password="pass")
        cls.branch = Branch.objects.create(code="BR-S", name_ar="فرع S")
        cls.device = QrDevice.objects.create(
            device_code="DEV-R1", branch=cls.branch, api_key_hash="hash", status="active",
        )
        cls.emp = Employee.objects.create(
            employee_code="S-001", first_name_ar="سعد", last_name_ar="س",
            branch=cls.branch, employment_status=Employee.EmploymentStatus.ACTIVE,
        )
        cls.day = AttendanceDay.objects.create(
            employee=cls.emp, work_date=datetime.date(2026, 8, 9),
            state=AttendanceDay.State.PRESENT,
        )
        AttendanceScan.objects.create(
            attendanceday=cls.day, employee=cls.emp, device=cls.device,
            source=AttendanceScan.Source.FIXED_READER,
            decision=AttendanceScan.Decision.CHECK_IN,
        )
        AttendanceScan.objects.create(
            employee=cls.emp, device=cls.device,
            source=AttendanceScan.Source.FIXED_READER,
            decision=AttendanceScan.Decision.REJECTED, result_detail="مفتاح غير صالح",
        )
        AttendanceScan.objects.create(
            employee=cls.emp, device=cls.device,
            source=AttendanceScan.Source.FIXED_READER,
            decision=AttendanceScan.Decision.REJECTED, result_detail="QR منتهي",
        )

    def setUp(self):
        self.client.force_login(self.admin)

    def test_rep16_lists_scans(self):
        resp = self.client.get(reverse("reports:rep16"))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "دخول")
        self.assertContains(resp, "مرفوض")

    def test_rep16_filters_by_decision(self):
        resp = self.client.get(reverse("reports:rep16"), {"decision": "rejected"})
        self.assertContains(resp, "QR منتهي")
        self.assertContains(resp, "مفتاح غير صالح")
        self.assertNotContains(resp, "<td>دخول</td>")

    def test_rep17_groups_rejected(self):
        resp = self.client.get(reverse("reports:rep17"))
        self.assertEqual(resp.status_code, 200)
        # صفّان: واحد لكل (موظف/جهاز/سبب) بتكرار 1
        self.assertEqual(resp.content.count(b"<tr><td>S-001"), 2)
        self.assertContains(resp, "QR منتهي")
        self.assertContains(resp, "مفتاح غير صالح")

    def test_rep17_export_csv(self):
        resp = self.client.get(reverse("reports:export", args=["rep17", "csv"]))
        self.assertEqual(resp.status_code, 200)
        rows = list(csv.reader(io.StringIO(resp.content.decode("utf-8"))))
        self.assertEqual(rows[0], ["الرمز", "الاسم", "الجهاز", "السبب", "عدد المرات"])
        self.assertEqual(rows[1][0], "S-001")


class PayrollReportTests(TestCase):
    """REP-30 — كشف الرواتب (ملخص) + صلاحيات التصدير المالي."""

    @classmethod
    def setUpTestData(cls):
        from apps.employees.models import Contract
        from apps.payroll.models import PayElement
        from apps.payroll.services import generate_payrun

        Permission.objects.create(code="reports.view", module="core", name_ar="التقارير")
        Permission.objects.create(code="reports.export", module="core", name_ar="تصدير")
        Permission.objects.create(code="payroll.payslip.view", module="payroll",
                                  name_ar="كشف الرواتب")
        cls.admin = User.objects.create_superuser(username="pr_admin", password="pass")
        cls.finance = User.objects.create_user(username="pr_fin", password="pass")
        cls.peon = User.objects.create_user(username="pr_peon", password="pass")
        cls.branch = Branch.objects.create(code="BR-PR", name_ar="فرع PR")
        cls.emp = Employee.objects.create(
            employee_code="PR-001", first_name_ar="مالي", last_name_ar="1",
            branch=cls.branch, employment_status=Employee.EmploymentStatus.ACTIVE,
            hire_date=datetime.date(2023, 1, 1), bank_account="ACC-PR",
        )
        Contract.objects.create(
            contract_number="CTR-PR", employee=cls.emp,
            start_date=datetime.date(2023, 1, 1), end_date=None,
            base_salary=Decimal("50000"), gross_salary=Decimal("50000"), allowance=0,
        )
        PayElement.objects.create(code="BASIC-PR", name_ar="أساسي", kind="earning",
                                  calculation="fixed", amount=Decimal("0"),
                                  applies_to_all=True, is_active=True)
        cls.payrun = generate_payrun("2026-08", cls.branch, cls.admin)

    def setUp(self):
        self.client.force_login(self.admin)

    def test_rep30_shows_summary(self):
        resp = self.client.get(reverse("reports:rep30"))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "2026-08")
        self.assertContains(resp, "50000.00")

    def test_rep30_requires_payroll_perm(self):
        self.client.force_login(self.peon)
        self.assertEqual(self.client.get(reverse("reports:rep30")).status_code, 403)

    def test_rep30_export_financial_guard(self):
        self.client.force_login(self.peon)
        resp = self.client.get(reverse("reports:export", args=["rep30", "csv"]))
        self.assertEqual(resp.status_code, 403)
        self.client.force_login(self.admin)
        resp = self.client.get(reverse("reports:export", args=["rep30", "csv"]))
        self.assertEqual(resp.status_code, 200)
        rows = list(csv.reader(io.StringIO(resp.content.decode("utf-8"))))
        self.assertEqual(rows[0][0], "الفترة")
        self.assertEqual(rows[1][0], "2026-08")

    def test_rep30_filter_by_period(self):
        resp = self.client.get(reverse("reports:rep30"), {"period_code": "2026-08"})
        self.assertContains(resp, "2026-08")
        resp = self.client.get(reverse("reports:rep30"), {"period_code": "1999-01"})
        self.assertContains(resp, "لا بيانات مطابقة")
        self.assertNotContains(resp, "<td>2026-08</td>")


class ReportV2Tests(TestCase):
    """REP-04/31/33/42/43/50/51/53 + حماية التقارير الحساسة."""

    @classmethod
    def setUpTestData(cls):
        Permission.objects.get_or_create(code="reports.view", defaults={"module": "core", "name_ar": "التقارير"})
        Permission.objects.get_or_create(code="payroll.payslip.view", defaults={"module": "payroll", "name_ar": "قسائم"})
        Permission.objects.get_or_create(code="system.audit.view", defaults={"module": "core", "name_ar": "تدقيق"})
        Permission.objects.get_or_create(code="system.backup.manage", defaults={"module": "core", "name_ar": "نسخ"})
        Permission.objects.get_or_create(code="device.manage", defaults={"module": "devices", "name_ar": "أجهزة"})
        cls.admin = User.objects.create_superuser(username="repv2_admin", password="pass")
        cls.peon = User.objects.create_user(username="repv2_peon", password="pass")
        cls.branch = Branch.objects.create(code="BR-V2", name_ar="فرع V2")
        cls.dept = Department.objects.create(code="DEP-V2", name_ar="قسم V2", branch=cls.branch)
        cls.emp = Employee.objects.create(
            employee_code="V2-001", first_name_ar="فائز", last_name_ar="V",
            branch=cls.branch, department=cls.dept, hire_date=datetime.date(2023, 1, 1),
        )
        from apps.employees.models import Contract

        cls.contract = Contract.objects.create(
            contract_number="CTR-V2", employee=cls.emp,
            contract_type=Contract.ContractType.CDD,
            start_date=datetime.date(2026, 1, 1),
            end_date=datetime.date.today() + datetime.timedelta(days=10),
            gross_salary=Decimal("50000"), base_salary=Decimal("45000"),
        )
        from apps.payroll.models import PayRun, Payslip, PayElement, PayrollLine, EndOfService

        cls.payrun = PayRun.objects.create(period_code="2026-08", branch=cls.branch,
                                           status=PayRun.Status.FROZEN)
        Payslip.objects.create(pay_run=cls.payrun, employee=cls.emp,
                               basic_salary=Decimal("45000"),
                               total_earnings=Decimal("50000"),
                               total_deductions=Decimal("5000"),
                               net=Decimal("45000"))
        cls.bonus = PayElement.objects.create(
            code="V2-BONUS", name_ar="مكافأة", kind=PayElement.Kind.EARNING,
            calculation=PayElement.Calculation.FIXED, amount=Decimal("5000"),
        )
        cls.deduct = PayElement.objects.create(
            code="V2-DED", name_ar="خصم غياب", kind=PayElement.Kind.DEDUCTION,
            calculation=PayElement.Calculation.FIXED, amount=Decimal("1000"),
        )
        PayrollLine.objects.create(pay_run=cls.payrun, employee=cls.emp,
                                   element=cls.bonus, amount=Decimal("5000"))
        PayrollLine.objects.create(pay_run=cls.payrun, employee=cls.emp,
                                   element=cls.deduct, amount=Decimal("1000"))
        EndOfService.objects.create(
            employee=cls.emp, termination_date=datetime.date(2026, 7, 31),
            total_years=Decimal("3.5"), service_reward=Decimal("40000"),
            unused_leave_comp=Decimal("5000"), notice_period=Decimal("0"),
            deductions=Decimal("2000"), net=Decimal("43000"),
            status=EndOfService.Status.APPROVED,
        )
        from apps.perf.models import PerfCycle, PerfReview, PipPlan

        cls.cycle = PerfCycle.objects.create(
            name_ar="دورة V2", period_start=datetime.date(2026, 1, 1),
            period_end=datetime.date(2026, 6, 30),
        )
        cls.review = PerfReview.objects.create(
            employee=cls.emp, cycle=cls.cycle,
            self_score=Decimal("80"), manager_score=Decimal("90"),
            final_score=Decimal("85"), status=PerfReview.Status.DONE,
        )
        PipPlan.objects.create(
            review=cls.review, start_date=datetime.date(2026, 7, 1),
            end_date=datetime.date(2026, 8, 1), action_items="تدريب",
        )
        from apps.backup.models import BackupJob

        BackupJob.objects.create(kind=BackupJob.Kind.DAILY, status=BackupJob.Status.SUCCESS)
        from apps.core.models import AuditLog

        AuditLog.objects.create(action=AuditLog.Action.CREATE, model_name="employee",
                                object_repr="V2-001", detail="اختبار")
        AuditLog.objects.create(action=AuditLog.Action.LOGIN, model_name="auth",
                                object_repr="repv2_admin", ip="10.0.0.1")
        from apps.devices.models import QrDevice
        from apps.attendance.models import AttendanceScan

        cls.device = QrDevice.objects.create(
            device_code="QR-V2", branch=cls.branch, api_key_hash="h", status=QrDevice.Status.ACTIVE,
        )
        AttendanceScan.objects.create(
            employee=cls.emp, device=cls.device, decision=AttendanceScan.Decision.CHECK_IN,
        )

    def setUp(self):
        self.client.force_login(self.admin)

    def test_rep04_lists_expiring_contracts(self):
        resp = self.client.get(reverse("reports:rep04"))
        self.assertContains(resp, "V2-001")

    def test_rep31_detail(self):
        resp = self.client.get(reverse("reports:rep31"))
        self.assertContains(resp, "V2-001")
        self.assertContains(resp, "45000")

    def test_rep33_cost(self):
        resp = self.client.get(reverse("reports:rep33"))
        self.assertContains(resp, "فرع V2")
        self.assertContains(resp, "45000")

    def test_rep34_bonuses_deductions(self):
        resp = self.client.get(reverse("reports:rep34"))
        self.assertContains(resp, "مكافأة")
        self.assertContains(resp, "خصم غياب")
        resp = self.client.get(reverse("reports:rep34"), {"kind": "earning"})
        self.assertContains(resp, "مكافأة")
        self.assertNotContains(resp, "خصم غياب")

    def test_rep35_end_of_service(self):
        resp = self.client.get(reverse("reports:rep35"))
        self.assertContains(resp, "V2-001")
        self.assertContains(resp, "43000")
        resp = self.client.get(reverse("reports:rep35"), {"status": "approved"})
        self.assertContains(resp, "V2-001")

    def test_rep52_daily_activity(self):
        resp = self.client.get(reverse("reports:rep52"))
        self.assertContains(resp, "repv2_admin")
        self.assertContains(resp, "10.0.0.1")

    def test_rep54_device_status(self):
        resp = self.client.get(reverse("reports:rep54"))
        self.assertContains(resp, "QR-V2")
        self.assertContains(resp, "فرع V2")

    def test_rep42_perf_results(self):
        resp = self.client.get(reverse("reports:rep42"))
        self.assertContains(resp, "85")

    def test_rep43_pip(self):
        resp = self.client.get(reverse("reports:rep43"))
        self.assertContains(resp, "V2-001")

    def test_rep50_audit(self):
        resp = self.client.get(reverse("reports:rep50"))
        self.assertContains(resp, "اختبار")

    def test_rep51_users(self):
        resp = self.client.get(reverse("reports:rep51"))
        self.assertContains(resp, "repv2_admin")

    def test_rep53_backup(self):
        resp = self.client.get(reverse("reports:rep53"))
        self.assertContains(resp, "يومي")

    def test_financial_guards(self):
        self.client.force_login(self.peon)
        self.assertEqual(self.client.get(reverse("reports:rep31")).status_code, 403)
        self.assertEqual(self.client.get(reverse("reports:rep33")).status_code, 403)
        self.assertEqual(self.client.get(reverse("reports:rep34")).status_code, 403)
        self.assertEqual(self.client.get(reverse("reports:rep35")).status_code, 403)
        self.assertEqual(self.client.get(reverse("reports:rep50")).status_code, 403)
        self.assertEqual(self.client.get(reverse("reports:rep52")).status_code, 403)
        self.assertEqual(self.client.get(reverse("reports:rep53")).status_code, 403)
        self.assertEqual(self.client.get(reverse("reports:rep54")).status_code, 403)

    def test_rep42_export_csv(self):
        resp = self.client.get(reverse("reports:export", args=["rep42", "csv"]))
        self.assertEqual(resp.status_code, 200)
        rows = list(csv.reader(io.StringIO(resp.content.decode("utf-8"))))
        self.assertEqual(rows[0][0], "الموظف")
        self.assertIn("V2-001", rows[1][0])


class ScheduledReportTests(TestCase):
    """التقارير المجدولة (T-REP-5): التنفيذ + الإشعارات + التنزيل."""

    @classmethod
    def setUpTestData(cls):
        Permission.objects.get_or_create(code="reports.view", defaults={"module": "core", "name_ar": "التقارير"})
        Permission.objects.get_or_create(code="reports.export", defaults={"module": "core", "name_ar": "تصدير"})
        cls.admin = User.objects.create_superuser(username="sch_admin", password="pass")
        cls.recipient = User.objects.create_user(username="sch_rec", password="pass")
        cls.branch = Branch.objects.create(code="BR-S", name_ar="فرع S")
        cls.dept = Department.objects.create(code="DEP-S", name_ar="قسم S", branch=cls.branch)
        cls.emp = Employee.objects.create(
            employee_code="S-001", first_name_ar="سامي", last_name_ar="S",
            branch=cls.branch, department=cls.dept,
        )
        from apps.reports.models import ReportDefinition

        cls.definition = ReportDefinition.objects.create(
            code="REP-01", name_ar="قائمة الموظفين", owner=cls.admin,
            template_type=ReportDefinition.Template.CSV, schedule=ReportDefinition.Schedule.DAILY,
            run_at=datetime.time(8, 30), filters_json={},
        )
        cls.definition.notify_users.add(cls.recipient)

    def setUp(self):
        from apps.reports.models import ReportDefinition

        now = datetime.datetime.now().replace(hour=8, minute=30)
        ReportDefinition.objects.filter(pk=self.definition.pk).update(last_run_at=None)
        self._now = now

    def test_is_due_daily(self):
        from apps.reports.scheduling import _is_due
        from apps.reports.models import ReportDefinition

        due = _is_due(self.definition, datetime.datetime.now().replace(hour=9, minute=0))
        self.assertTrue(due)
        not_due = _is_due(self.definition, datetime.datetime.now().replace(hour=8, minute=0))
        self.assertFalse(not_due)

    def test_run_creates_job_and_file(self):
        from apps.notif.models import Notification
        from apps.reports.models import ReportJob

        jobs = run_scheduled_reports()
        self.assertEqual(len(jobs), 1)
        job = jobs[0]
        self.assertEqual(job.status, ReportJob.Status.DONE)
        self.assertEqual(job.files.count(), 1)
        f = job.files.first()
        self.assertEqual(f.format, "csv")
        self.assertGreater(f.size_bytes, 0)
        self.assertTrue(Path(f.file_path).exists())
        self.assertTrue(Notification.objects.filter(user=self.recipient).exists())

    def test_no_rerun_same_window(self):
        from apps.reports.models import ReportJob

        self.assertEqual(len(run_scheduled_reports()), 1)
        self.assertEqual(len(run_scheduled_reports()), 0)
        self.assertEqual(ReportJob.objects.count(), 1)

    def test_invalid_code_fails_job(self):
        from apps.reports.models import ReportDefinition, ReportJob

        ReportDefinition.objects.create(
            code="REP-99", name_ar="غير معروف", owner=self.admin,
            template_type=ReportDefinition.Template.CSV, schedule=ReportDefinition.Schedule.DAILY,
            run_at=datetime.time(8, 30),
        )
        jobs = run_scheduled_reports()
        failed = [j for j in jobs if j.status == ReportJob.Status.FAILED]
        self.assertEqual(len(failed), 1)

    def test_generated_list_and_download(self):
        self.client.force_login(self.admin)
        run_scheduled_reports()
        resp = self.client.get(reverse("reports:generated"))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "قائمة الموظفين")
        f = self.definition.jobs.first().files.first()
        resp = self.client.get(reverse("reports:generated_download", args=[f.pk]))
        self.assertEqual(resp.status_code, 200)

    def test_download_guard(self):
        self.client.force_login(self.recipient)
        self.assertEqual(self.client.get(reverse("reports:generated")).status_code, 403)


class OperationalReportsV2Tests(TestCase):
    """REP-02/03/14/15/18/22/23 — التقارير التشغيلية المكملة."""

    @classmethod
    def setUpTestData(cls):
        Permission.objects.get_or_create(code="reports.view", defaults={"module": "core", "name_ar": "التقارير"})
        cls.admin = User.objects.create_superuser(username="op_admin", password="pass")
        cls.branch = Branch.objects.create(code="BR-O", name_ar="فرع O")
        cls.dept = Department.objects.create(code="DEP-O", name_ar="قسم O", branch=cls.branch)
        from apps.org.models import Position

        cls.pos = Position.objects.create(code="POS-O", name_ar="مهندس", department=cls.dept)
        cls.emp = Employee.objects.create(
            employee_code="O-001", first_name_ar="عمر", last_name_ar="O",
            branch=cls.branch, department=cls.dept, position=cls.pos,
            hire_date=datetime.date(2026, 7, 1),
        )
        cls.today = datetime.date(2026, 8, 9)
        AttendanceDay.objects.create(employee=cls.emp, work_date=cls.today,
                                     state=AttendanceDay.State.PRESENT,
                                     early_minutes=20, overtime_minutes=90,
                                     worked_minutes=480)
        from apps.attendance.models import AttendanceException
        from apps.leave.models import PublicHoliday

        cls.lt = LeaveType.objects.create(code="op-lt", name_ar="سنوية", days_per_year=30)
        AttendanceException.objects.create(
            employee=cls.emp, type=AttendanceException.Type.MISSION,
            from_time=datetime.datetime(2026, 8, 9, 10, 0),
            to_time=datetime.datetime(2026, 8, 9, 12, 0), hours=2,
            status=AttendanceException.Status.APPROVED, approved_by=cls.admin,
        )
        LeaveRequest.objects.create(employee=cls.emp, leave_type=cls.lt,
                                    from_date=datetime.date(2026, 8, 1),
                                    to_date=datetime.date(2026, 8, 31), days=5,
                                    status=LeaveRequest.Status.APPROVED)
        PublicHoliday.objects.create(branch=cls.branch, date=datetime.date(2026, 7, 5),
                                     name_ar="عيد الفطر", is_recurring=True)

    def setUp(self):
        self.client.force_login(self.admin)

    def test_rep02_org_structure(self):
        resp = self.client.get(reverse("reports:rep02"))
        self.assertContains(resp, "مهندس")
        self.assertContains(resp, "قسم O")

    def test_rep03_new_employees(self):
        resp = self.client.get(reverse("reports:rep03"),
                               {"from_date": "2026-01-01", "to_date": "2026-12-31"})
        self.assertContains(resp, "O-001")
        resp = self.client.get(reverse("reports:rep03"), {"from_date": "2027-01-01"})
        self.assertNotContains(resp, "O-001")

    def test_rep14_early_departure(self):
        resp = self.client.get(reverse("reports:rep14"),
                               {"from_date": "2026-08-01", "to_date": "2026-08-31"})
        self.assertContains(resp, "20")

    def test_rep15_overtime(self):
        resp = self.client.get(reverse("reports:rep15"),
                               {"from_date": "2026-08-01", "to_date": "2026-08-31"})
        self.assertContains(resp, "90")

    def test_rep18_exceptions(self):
        resp = self.client.get(reverse("reports:rep18"))
        self.assertContains(resp, "مأمورية")
        self.assertContains(resp, "op_admin")

    def test_rep22_ongoing_leaves(self):
        resp = self.client.get(reverse("reports:rep22"), {"work_date": "2026-08-15"})
        self.assertContains(resp, "O-001")
        resp = self.client.get(reverse("reports:rep22"), {"work_date": "2026-12-15"})
        self.assertNotContains(resp, "O-001")

    def test_rep23_public_holidays(self):
        resp = self.client.get(reverse("reports:rep23"), {"year": "2026"})
        self.assertContains(resp, "عيد الفطر")


class ExecutiveReportsV2Tests(TestCase):
    """REP-05/06/60/61/62 — التقاعد وتغييرات الوظائف والتنفيذية."""

    @classmethod
    def setUpTestData(cls):
        Permission.objects.get_or_create(code="reports.view", defaults={"module": "core", "name_ar": "التقارير"})
        cls.admin = User.objects.create_superuser(username="ex_admin", password="pass")
        cls.branch = Branch.objects.create(code="BR-X", name_ar="فرع X")
        cls.dept = Department.objects.create(code="DEP-X", name_ar="قسم X", branch=cls.branch)
        from apps.org.models import Position

        cls.pos1 = Position.objects.create(code="POS-X1", name_ar="منفذ", department=cls.dept)
        cls.pos2 = Position.objects.create(code="POS-X2", name_ar="مشرف", department=cls.dept)
        cls.emp = Employee.objects.create(
            employee_code="X-001", first_name_ar="خالد", last_name_ar="X",
            branch=cls.branch, department=cls.dept, position=cls.pos2,
            birth_date=datetime.date(1966, 6, 15), hire_date=datetime.date(2026, 6, 1),
        )
        cls.today = datetime.date(2026, 8, 9)
        from apps.employees.models import EmploymentHistory

        EmploymentHistory.objects.create(
            employee=cls.emp, branch=cls.branch, department=cls.dept, position=cls.pos1,
            effective_from=datetime.date(2026, 6, 1), is_current=False,
        )
        EmploymentHistory.objects.create(
            employee=cls.emp, branch=cls.branch, department=cls.dept, position=cls.pos2,
            effective_from=datetime.date(2026, 7, 15), is_current=True,
        )
        from apps.payroll.models import EndOfService, PayRun, Payslip

        cls.payrun = PayRun.objects.create(period_code="2026-07", branch=cls.branch,
                                           status=PayRun.Status.FROZEN)
        Payslip.objects.create(pay_run=cls.payrun, employee=cls.emp,
                               basic_salary=Decimal("45000"),
                               total_earnings=Decimal("45000"),
                               total_deductions=Decimal("0"), net=Decimal("45000"))
        AttendanceDay.objects.create(employee=cls.emp, work_date=cls.today,
                                     state=AttendanceDay.State.ABSENT, worked_minutes=0)
        AttendanceDay.objects.create(employee=cls.emp,
                                     work_date=datetime.date(2026, 8, 10),
                                     state=AttendanceDay.State.PRESENT, worked_minutes=480)
        EndOfService.objects.create(employee=cls.emp, termination_date=datetime.date(2026, 8, 5),
                                    net=Decimal("30000"), status=EndOfService.Status.PAID)

    def setUp(self):
        self.client.force_login(self.admin)

    def test_rep05_retirement(self):
        resp = self.client.get(reverse("reports:rep05"),
                               {"from_date": "2026-01-01", "to_date": "2026-12-31"})
        self.assertContains(resp, "X-001")
        resp = self.client.get(reverse("reports:rep05"),
                               {"from_date": "2030-01-01", "to_date": "2035-01-01"})
        self.assertNotContains(resp, "X-001")

    def test_rep06_job_changes(self):
        resp = self.client.get(reverse("reports:rep06"))
        self.assertContains(resp, "منفذ")
        self.assertContains(resp, "مشرف")
        resp = self.client.get(reverse("reports:rep06"), {"change_type": "promotion"})
        self.assertContains(resp, "ترقية/منصب")

    def test_rep60_executive_summary(self):
        resp = self.client.get(reverse("reports:rep60"),
                               {"from_date": "2026-08-01", "to_date": "2026-08-31"})
        self.assertContains(resp, "فرع X")
        self.assertContains(resp, "45000")

    def test_rep61_kpi(self):
        resp = self.client.get(reverse("reports:rep61"),
                               {"from_date": "2026-08-01", "to_date": "2026-08-31"})
        self.assertContains(resp, "نسبة الحضور")
        self.assertContains(resp, "معدل الدوران")

    def test_rep62_turnover(self):
        resp = self.client.get(reverse("reports:rep62"),
                               {"from_date": "2026-08-01", "to_date": "2026-08-31"})
        self.assertContains(resp, "قسم X")


class TrilingualReportTests(TestCase):
    """التقارير بثلاث لغات (ar/fr/en) — مصطلحات وعناوين وبيانات مناسبة لكل لغة."""

    @classmethod
    def setUpTestData(cls):
        Permission.objects.create(code="reports.view", module="core", name_ar="التقارير")
        Permission.objects.create(code="reports.export", module="core", name_ar="تصدير")
        cls.admin = User.objects.create_superuser(username="tri_admin", password="pass")
        cls.branch = Branch.objects.create(
            code="BR-T", name_ar="فرع الجزائر", name_fr="Succursale Alger",
            name_en="Algiers Branch",
        )
        cls.dept = Department.objects.create(
            code="DEP-T", name_ar="قسم الموارد", name_fr="Département RH",
            name_en="HR Department", branch=cls.branch,
        )
        cls.emp = Employee.objects.create(
            employee_code="T-001",
            first_name_ar="أمين", last_name_ar="بن يوسف",
            first_name_fr="Amine", last_name_fr="Ben Youcef",
            first_name_en="Amine", last_name_en="Ben Youcef",
            branch=cls.branch, department=cls.dept,
            employment_status=Employee.EmploymentStatus.ACTIVE, phone="0550",
        )

    def setUp(self):
        self.client.force_login(self.admin)

    def test_rep01_arabic_page(self):
        resp = self.client.get(reverse("reports:rep01"), {"lang": "ar"})
        self.assertContains(resp, "قائمة الموظفين")
        self.assertContains(resp, "أمين بن يوسف")
        self.assertContains(resp, "قسم الموارد")

    def test_rep01_french_page(self):
        resp = self.client.get(reverse("reports:rep01"), {"lang": "fr"})
        self.assertContains(resp, "Liste des employés")
        self.assertContains(resp, "Département")
        self.assertContains(resp, "Amine Ben Youcef")
        self.assertContains(resp, "Département RH")

    def test_rep01_english_page(self):
        resp = self.client.get(reverse("reports:rep01"), {"lang": "en"})
        self.assertContains(resp, "Employee list")
        self.assertContains(resp, "Department")
        self.assertContains(resp, "Algiers Branch")

    def test_rep01_language_selector_present(self):
        resp = self.client.get(reverse("reports:rep01"), {"lang": "fr"})
        self.assertContains(resp, "lang=ar")
        self.assertContains(resp, "lang=en")
        self.assertContains(resp, "Français")

    def test_export_csv_french(self):
        resp = self.client.get(reverse("reports:export", args=["rep01", "csv"]), {"lang": "fr"})
        self.assertEqual(resp.status_code, 200)
        rows = list(csv.reader(io.StringIO(resp.content.decode("utf-8"))))
        self.assertEqual(rows[0], ["Code", "Nom", "Département", "Poste", "Succursale", "Statut", "Téléphone"])
        self.assertIn("Amine Ben Youcef", rows[1])
        self.assertIn("Département RH", rows[1])

    def test_export_csv_english(self):
        resp = self.client.get(reverse("reports:export", args=["rep01", "csv"]), {"lang": "en"})
        self.assertEqual(resp.status_code, 200)
        rows = list(csv.reader(io.StringIO(resp.content.decode("utf-8"))))
        self.assertEqual(rows[0], ["Code", "Name", "Department", "Position", "Branch", "Status", "Phone"])
        self.assertIn("Algiers Branch", rows[1])

    def test_export_csv_arabic(self):
        resp = self.client.get(reverse("reports:export", args=["rep01", "csv"]), {"lang": "ar"})
        rows = list(csv.reader(io.StringIO(resp.content.decode("utf-8"))))
        self.assertEqual(rows[0][0], "الرمز")
        self.assertIn("أمين بن يوسف", rows[1])

    def test_export_xlsx_french(self):
        from openpyxl import load_workbook

        resp = self.client.get(reverse("reports:export", args=["rep01", "xlsx"]), {"lang": "fr"})
        self.assertEqual(resp.status_code, 200)
        wb = load_workbook(io.BytesIO(resp.content))
        ws = wb.active
        self.assertEqual(ws["A1"].value, "Code")
        self.assertEqual(ws["C1"].value, "Département")
        self.assertTrue(any("Amine Ben Youcef" in str(c.value) for row in ws.iter_rows() for c in row if c.value))

    def test_export_pdf_all_languages(self):
        for lang in ("ar", "fr", "en"):
            with self.subTest(lang=lang):
                resp = self.client.get(reverse("reports:export", args=["rep01", "pdf"]), {"lang": lang})
                self.assertEqual(resp.status_code, 200)
                self.assertEqual(resp["Content-Type"], "application/pdf")
                self.assertTrue(resp.content.startswith(b"%PDF"))

    def test_pdf_template_direction(self):
        from django.template.loader import render_to_string

        for lang, direction, title in (("ar", "rtl", "قائمة الموظفين"),
                                       ("fr", "ltr", "Liste des employés"),
                                       ("en", "ltr", "Employee list")):
            with self.subTest(lang=lang):
                html = render_to_string("reports/pdf_report.html", {
                    "title": title, "header": ["X"], "rows": [["1"]],
                    "generated_at": "2026-08-09", "report_lang": lang,
                    "direction": direction, "body_font": "sans-serif", "page_font": "sans-serif",
                })
                self.assertIn(f'dir="{direction}"', html)
                self.assertIn(f'lang="{lang}"', html)

    def test_triname_filter(self):
        from django.template import Context, Template
        from django.utils import translation

        with translation.override("fr"):
            out = Template("{% load report_local %}{{ b|triname }}").render(Context({"b": self.branch}))
            self.assertEqual(out, "Succursale Alger")
        with translation.override("en"):
            out = Template("{% load report_local %}{{ b|triname }}").render(Context({"b": self.branch}))
            self.assertEqual(out, "Algiers Branch")
        with translation.override("ar"):
            out = Template("{% load report_local %}{{ b|triname }}").render(Context({"b": self.branch}))
            self.assertEqual(out, "فرع الجزائر")

    def test_french_status_terminology(self):
        resp = self.client.get(reverse("reports:rep01"), {"lang": "fr"})
        self.assertContains(resp, "Actif")
        resp = self.client.get(reverse("reports:export", args=["rep01", "csv"]), {"lang": "fr"})
        rows = list(csv.reader(io.StringIO(resp.content.decode("utf-8"))))
        self.assertIn("Actif", rows[1])
